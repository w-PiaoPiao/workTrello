# -*- coding: utf-8 -*-
"""
资金分配监视泳道导出测试：行组装（三线+配合、环节并集列序、日期归位）
与 Excel 写入（表头/冻结/读回验证）
"""

import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["PET_BOARD_DATA_DIR"] = tempfile.mkdtemp()

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from PySide6.QtWidgets import QApplication

_qapp = QApplication.instance() or QApplication([])

from app.models.board import BoardList, Card, FUND_STEP_TEMPLATES
from app.services.fund_export import (
    FUND_EXPORT_STEPS,
    build_rows,
    export_xlsx,
)


def _make_list() -> BoardList:
    lst = BoardList(title="资金分配监视", fund_watch=True)

    regular = Card(title="常务会项目", notes="第一轮分配")
    regular.fund = {"role": "lead", "kind": "regular",
                    "start_date": "2026-09-01", "steps": []}
    regular.switch_fund_kind("regular")
    regular.fund["start_date"] = "2026-09-01"
    for i in (0, 1):   # 收集分配方案、汇总上会材料已办
        regular.fund["steps"][i].update(
            done=True, date=f"2026-09-0{4 + i * 2}")

    branch = Card(title="分管项目")
    branch.switch_fund_kind("branch")
    branch.fund["role"] = "lead"
    branch.fund["steps"][8].update(done=True, date="2026-10-12")  # 财政二次去函

    finance = Card(title="财政项目")
    finance.switch_fund_kind("finance")
    finance.fund["role"] = "lead"

    assist = Card(title="配合项目", notes="台账跟进")
    assist.ensure_fund()

    lst.cards = [regular, branch, finance, assist,
                 Card(title="无监视配置"), ]
    archived = Card(title="已归档", archived=True)
    archived.ensure_fund()
    lst.cards.append(archived)
    return lst


class BuildRowsTest(unittest.TestCase):
    def setUp(self):
        self.lst = _make_list()
        self.rows = build_rows(self.lst)

    def test_row_selection_and_count(self):
        """只导出带监视配置且未归档的卡片"""
        self.assertEqual([r[0] for r in self.rows],
                         ["常务会项目", "分管项目", "财政项目", "配合项目"])

    def test_column_layout(self):
        """列 = 标题/类型/备注 + 环节并集（19 列）"""
        self.assertEqual(len(FUND_EXPORT_STEPS), 19)
        self.assertEqual(FUND_EXPORT_STEPS[15], "财政二次去函")
        self.assertEqual(FUND_EXPORT_STEPS[-3:],
                         ["会签下达文件", "资金文件下达", "资金入库追加"])
        for row in self.rows:
            self.assertEqual(len(row), 3 + len(FUND_EXPORT_STEPS))

    def test_role_labels(self):
        self.assertEqual(self.rows[0][1], "牵头·常务会")
        self.assertEqual(self.rows[1][1], "牵头·分管")
        self.assertEqual(self.rows[2][1], "牵头·财政")
        self.assertEqual(self.rows[3][1], "配合分配")

    def test_dates_landed_in_union_columns(self):
        """环节办理日期按并集列序归位；卡片没有的环节留空"""
        regular = self.rows[0]
        idx = FUND_EXPORT_STEPS.index("收集分配方案")
        self.assertEqual(regular[3 + idx], "2026-09-04")
        self.assertEqual(regular[3 + idx + 1], "2026-09-06")
        self.assertEqual(regular[3 + FUND_EXPORT_STEPS.index("上会")], "")
        branch = self.rows[1]
        self.assertEqual(branch[3 + FUND_EXPORT_STEPS.index("财政二次去函")],
                         "2026-10-12")
        # 常务会卡没有"财政二次去函"环节 → 空
        self.assertEqual(regular[3 + FUND_EXPORT_STEPS.index("财政二次去函")], "")
        # 配合卡环节列全空
        self.assertTrue(all(v == "" for v in self.rows[3][3:]))

    def test_notes_exported(self):
        self.assertEqual(self.rows[0][2], "第一轮分配")
        self.assertEqual(self.rows[3][2], "台账跟进")

    def test_empty_list_gives_no_rows(self):
        lst = BoardList(title="空", fund_watch=True)
        self.assertEqual(build_rows(lst), [])


class ExportXlsxTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "export.xlsx"
        self.n = export_xlsx(_make_list(), self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_written_and_read_back(self):
        """写出的 xlsx 可读回：表头/行数/单元格值/冻结窗格"""
        from openpyxl import load_workbook

        self.assertEqual(self.n, 4)
        wb = load_workbook(self.path)
        ws = wb.active
        header = [c.value for c in ws[1]]
        self.assertEqual(header[:3], ["卡片标题", "类型", "备注"])
        self.assertEqual(header[3:], FUND_EXPORT_STEPS)
        self.assertEqual(ws.cell(row=2, column=1).value, "常务会项目")
        self.assertEqual(ws.cell(row=2, column=2).value, "牵头·常务会")
        self.assertEqual(ws.max_row, 1 + 4)
        self.assertEqual(ws.max_column, 3 + len(FUND_EXPORT_STEPS))
        self.assertEqual(ws.freeze_panes, "B2")

    def test_headers_bold(self):
        from openpyxl import load_workbook

        ws = load_workbook(self.path).active
        self.assertTrue(all(c.font.bold for c in ws[1]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
